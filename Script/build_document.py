from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/private/tmp/hinton-ontology-doc.XiKwug')
OUTPUT = Path('/Users/peter/Ontology-Hinton/doc/大学新生本体类与属性设计.docx')

# Class name, Chinese label, direct superclass, meaning.
CLASS_GROUPS = [
('人员与组织', [
('Person','人员','owl:Thing','有独立身份的人。姓名等通用属性在此定义。'),
('Student','学生','Person','大学在读学生；按需要录入具体学生。'),
('UndergraduateStudent','本科生','Student','正在修读本科专业的学生。'),
('PostgraduateStudent','研究生','Student','正在修读研究生专业的学生。'),
('Staff','教职工','Person','承担教学、行政或支持工作的人。'),
('AcademicStaff','教学人员','Staff','担任授课、模块负责或专业负责工作的人。'),
('ProfessionalServicesStaff','行政与支持人员','Staff','提供教务、图书馆或技术等支持的人。'),
('Organisation','组织','owl:Thing','大学及其组织单位。'),
('University','大学','Organisation','大学机构本身，与校区地点区分。'),
('School','学院','Organisation','大学中的学院或相应学术单位。'),
('Department','系','Organisation','组织中的系级学术单位。'),
]),
('专业与教学安排', [
('Course','专业或学位项目','owl:Thing','完整培养项目。与旧文档 Programme 对应。'),
('UndergraduateCourse','本科专业','Course','通向本科学位的培养项目。'),
('PostgraduateCourse','研究生专业','Course','通向研究生学位的培养项目。'),
('Module','教学模块','owl:Thing','相对稳定的模块目录条目，记录内容和学分。'),
('ModuleOffering','模块开设','owl:Thing','一个模块在指定学年的具体开设，可跨多个学期。'),
('CourseModule','专业模块安排','owl:Thing','某专业某版培养方案中的模块条目，记录年级及必选属性。'),
('CourseLeadershipAppointment','专业负责人任职','owl:Thing','某位教学人员在一段期间担任某专业负责人的记录。'),
('Assessment','考核项目','owl:Thing','某次开课的具体考核组成项，记录权重和截止时间。'),
('Exam','考试','Assessment','以考试形式实施的考核项目。'),
('Coursework','课程作业','Assessment','以作业、项目或报告等形式实施的考核项目。'),
]),
('校园空间', [
('Place','地点','owl:Thing','可定位的物理空间。'),
('Campus','校区','Place','大学的一个校区，与 University 组织区分。'),
('Building','建筑','Place','校区内的一栋建筑。'),
('LibraryBuilding','图书馆建筑','Building','用于图书馆功能的建筑，不表示图书馆管理机构。'),
('Room','房间','Place','可独立识别或编号的室内房间。'),
('LectureTheatre','阶梯教室','Room','适合集中讲授的阶梯式教学房间。'),
('SeminarRoom','研讨室','Room','适合小组讨论与研讨的房间。'),
('Laboratory','实验室','Room','用于实验或实践教学的房间。'),
('ComputerLaboratory','机房','Laboratory','用于计算机实践教学的实验室。'),
('Office','办公室','Room','供教职工办公的房间。'),
('StudySpace','学习室','Room','本版限定为可独立编号的学习房间，可与其他用途类重叠。'),
('AccessibleRoom','无台阶可达房间','Room','定义类：Room 且 isStepFree 为 true。'),
]),
('活动时间与设备', [
('Event','活动','owl:Thing','一次有起止时间的活动。'),
('TeachingSession','教学场次','Event','一次具体授课安排，不表示整学期的模块。'),
('LectureSession','讲授场次','TeachingSession','一次讲授课。'),
('TutorialSession','辅导场次','TeachingSession','一次辅导或讨论课。'),
('PracticalSession','实践场次','TeachingSession','一次实验或实践课。'),
('OnCampusTeachingSession','校内房间授课场次','TeachingSession','本版指安排在校内一个房间的场次，可与讲授等类型重叠。'),
('CampusEvent','校园活动','Event','校园活动接口；具体子类随活动需求确定。'),
('TimePeriod','时间段','owl:Thing','学年、学期等有起止日期的时间段。'),
('AcademicYear','学年','TimePeriod','一个学年，如 AY2026_27。'),
('Semester','学期','TimePeriod','某个学年内的一个具体学期，不能跨年复用同一实例。'),
('EquipmentType','设备类型','owl:Thing','设备种类的分类概念，例如电脑类型或投影仪类型。'),
('EquipmentProvision','设备配置记录','owl:Thing','某个房间配置某类设备及其数量的一条记录。'),
]),
('设施匹配扩展', [
('FacilityType','设施类型','owl:Thing','Q21 使用的受控设施类别，例如实验设施、计算设施。'),
('FacilityRequirement','设施要求','owl:Thing','一项设施要求，可列出多个可接受的替代设施类型。'),
])]

# Property, Domain, Range, intended meaning / characteristics.
OBJECT_GROUPS = [
('人员组织与专业', [
('memberOf','Staff','Organisation','人员所属学院或系；允许多个组织。'),
('organisationPartOf','Organisation','Organisation','直接组织隶属；不默认传递，不与空间包含混用。'),
('offeredBy','Course','Department','开设专业的系；联合开设可有多个值。'),
('hasOffice','Staff','Office','当前办公地点；可有多个办公室。'),
('hasLeadershipAppointment','Course','CourseLeadershipAppointment','专业关联带任期的负责人任职记录。'),
('appointedStaffMember','CourseLeadershipAppointment','AcademicStaff','该条任职记录对应的教学人员。'),
('hasCourseModule','Course','CourseModule','专业关联培养方案中的模块条目。'),
('forCourse','CourseModule','Course','条目所属专业；hasCourseModule 的逆属性。'),
('forModule','CourseModule','Module','条目安排的教学模块。'),
]),
('开课与考核', [
('offeringOf','ModuleOffering','Module','具体开课对应的模块目录条目。'),
('inAcademicYear','ModuleOffering or CourseModule or Semester','AcademicYear','开课、培养条目或学期适用的学年；Domain 使用单个并集表达式。'),
('inSemester','ModuleOffering','Semester','具体开课覆盖的学期；允许全年模块连接多个学期。'),
('taughtBy','ModuleOffering','AcademicStaff','本次开课的教学团队；允许多人。'),
('hasModuleLeader','ModuleOffering','AcademicStaff','本次开课的负责人；允许联合负责。'),
('hasTeachingSession','ModuleOffering','TeachingSession','开课包含的具体场次；合班场次可被多个开课引用。'),
('ledBy','TeachingSession','AcademicStaff','该场次实际授课教师；不等同于全部教学团队。'),
('hasAssessment','ModuleOffering','Assessment','本次开课的考核组成项。'),
('hasDirectPrerequisite','Module','Module','直接先修要求；是 hasPrerequisite 的子属性。'),
('hasPrerequisite','Module','Module','直接或间接先修依赖；设为 Transitive。'),
('hasCorequisite','Module','Module','该模块要求同时或按规定修读的共修模块；不默认对称。'),
]),
('地点与设备', [
('takesPlaceIn','Event','Place','活动发生地点；校内房间授课另用类限制约束为 Room。'),
('locatedInBuilding','Room','Building','房间直接位于的建筑。'),
('locatedOnCampus','Building','Campus','建筑位于的校区。'),
('withinCampus','Room','Campus','由房间到建筑、建筑到校区的属性链推得。'),
('hasStudySpace','LibraryBuilding','StudySpace','图书馆包含的学习室；为 inverse locatedInBuilding 的子属性。'),
('hasEquipmentProvision','Room','EquipmentProvision','房间关联设备配置记录。'),
('equipmentType','EquipmentProvision','EquipmentType','配置记录对应的设备种类。'),
]),
('校园活动与设施扩展', [
('organisedBy','CampusEvent','Organisation','校园活动的组织方；初版接口。'),
('hasFacilityRequirement','Module or TeachingSession','FacilityRequirement','模块描述要求；需核验的场次显式关联所适用的要求记录。'),
('acceptableFacilityType','FacilityRequirement','FacilityType','同一要求记录下的多个类型表示任选其一。'),
('providesFacilityType','Room','FacilityType','房间已核实提供的设施类型；由可信数据或明确映射赋值。'),
])]

# Property, Domain, datatype, description.
DATA_GROUPS = [
('人员与组织', [
('fullName','Person','xsd:string','姓名。'),
('emailAddress','Person','xsd:string','联系邮箱；教师使用官方邮箱。'),
('jobTitle','Staff','xsd:string','职称或岗位；无需为每个职称建立子类。'),
('telephoneNumber','Staff','xsd:string','电话号码；保留区号和前导零。'),
('staffProfileURL','Staff','xsd:anyURI','官方人员介绍页面。'),
('staffID','Staff','xsd:string','校内人员标识，可选。'),
('studentID','Student','xsd:string','学生标识；仅在有个人化需求时使用。'),
('organisationName','Organisation','xsd:string','组织名称。'),
('organisationCode','Organisation','xsd:string','组织代码，可选。'),
]),
('专业模块与任职', [
('courseTitle','Course','xsd:string','专业名称。'),
('courseCode','Course','xsd:string','专业代码。'),
('courseDescription','Course','xsd:string','专业简介。'),
('qualificationName','Course','xsd:string','授予资格名称，如 BSc、MSc。'),
('durationYears','Course','xsd:decimal','标准学制，单位为年。'),
('studyMode','Course','xsd:string','全日制或非全日制；不同学制路径需分别描述。'),
('moduleTitle','Module','xsd:string','模块标题。'),
('moduleCode','Module','xsd:string','模块代码。'),
('moduleDescription','Module','xsd:string','模块内容概述。'),
('learningOutcome','Module','xsd:string','学习目标；可有多个值。'),
('credits','Module','xsd:nonNegativeInteger','学分；本版假定模块目录中的学分稳定。'),
('academicLevel','Module','xsd:string','资格框架等级及其约定；不等同于学习年级。'),
('offeringCode','ModuleOffering','xsd:string','区分年度、班次或分组开课的标识。'),
('studyYear','CourseModule','xsd:positiveInteger','该模块安排在专业的第几年，如 1。'),
('isCompulsory','CourseModule','xsd:boolean','true 为必修，false 为选修；缺失表示未知。'),
('curriculumVersion','CourseModule','xsd:string','培养方案版本；查询时与学年一起选择。'),
('validFrom','CourseLeadershipAppointment','xsd:date','任期开始日期。'),
('validUntil','CourseLeadershipAppointment','xsd:date','任期最后一天；缺失不能自动解释为仍在任。'),
]),
('空间与预约', [
('campusName','Campus','xsd:string','校区名称。'),
('buildingName','Building','xsd:string','建筑名称。'),
('buildingCode','Building','xsd:string','建筑代码。'),
('postalAddress','Building','xsd:string','建筑地址。'),
('mapURL','Building','xsd:anyURI','官方地图链接。'),
('latitude','Building','xsd:decimal','纬度，范围为负 90 至 90。'),
('longitude','Building','xsd:decimal','经度，范围为负 180 至 180。'),
('hasLift','Building','xsd:boolean','是否配有电梯；不单独证明房间可无台阶到达。'),
('accessibleEntrance','Building','xsd:boolean','是否至少有一个已核实的无台阶入口。'),
('roomID','Room','xsd:string','跨建筑唯一的房间标识。'),
('roomNumber','Room','xsd:string','门牌号；本身不保证全校唯一。'),
('floorNumber','Room','xsd:string','楼层标记，允许 G、B1、8 等学校实际编码。'),
('capacity','Room','xsd:nonNegativeInteger','允许容纳人数；用于计算时须大于零。'),
('isStepFree','Room','xsd:boolean','是否已核实存在通往该房间的无台阶路线。'),
('isBookable','StudySpace','xsd:boolean','是否支持提前预约。'),
('bookingRequired','StudySpace','xsd:boolean','使用是否必须预约；与可预约不同。'),
('bookingURL','StudySpace','xsd:anyURI','预约页面；是否已被占用不由此属性表达。'),
]),
('设备与设施', [
('equipmentName','EquipmentType','xsd:string','设备类型名称，如电脑、投影仪。'),
('quantity','EquipmentProvision','xsd:nonNegativeInteger','该房间该类设备数量。'),
('facilityTypeName','FacilityType','xsd:string','设施类别名称，供 Q21 的受控词表使用。'),
('requirementDescription','FacilityRequirement','xsd:string','设施要求的人类可读解释。'),
]),
('场次考核与时间', [
('eventTitle','Event','xsd:string','教学场次或校园活动的名称。'),
('startDateTime','Event','xsd:dateTime','活动开始时刻；数据应带一致的时区偏移。'),
('endDateTime','Event','xsd:dateTime','活动结束时刻。'),
('eventStatus','Event','xsd:string','受控值 Scheduled、Cancelled、Completed。'),
('expectedStudentCount','TeachingSession','xsd:nonNegativeInteger','该场次预计到课人数；不使用整模块人数替代分组人数。'),
('assessmentTitle','Assessment','xsd:string','考核项目名称。'),
('weightingPercentage','Assessment','xsd:decimal','权重百分数，0 至 100；例如 40 表示 40%。'),
('dueDateTime','Assessment','xsd:dateTime','已公布的截止时刻，需包含时区。'),
('submissionInstructions','Assessment','xsd:string','提交要求。'),
('submissionURL','Assessment','xsd:anyURI','提交入口。'),
('periodName','TimePeriod','xsd:string','学年或学期名称。'),
('startDate','TimePeriod','xsd:date','时间段起始日期。'),
('endDate','TimePeriod','xsd:date','时间段最后一天，含当日。'),
('eventURL','CampusEvent','xsd:anyURI','活动信息页；暂定接口。'),
('registrationURL','CampusEvent','xsd:anyURI','活动报名地址；暂定接口。'),
('targetAudience','CampusEvent','xsd:string','目标参与人群说明；复杂资格条件后续单独建模。'),
('lastVerifiedOn','owl:Thing','xsd:date','当前实体信息最后核实日期；由维护者更新。'),
])]

ANNOTATIONS = [
('rdfs:label','类、属性、个体','中文与英文标签，使用 @zh 和 @en。'),
('rdfs:comment','类、属性、个体','定义、适用边界及必要说明。'),
('skos:altLabel','个体或术语','别名，如 Science Building 与 SCI Building。'),
('dcterms:source','个体或公理','来源页面或资料的 IRI；与 lastVerifiedOn 配合。'),
]

MAPPING = [
('Q1','模块教师与负责人','ModuleOffering、Module、AcademicStaff','offeringOf、taughtBy、hasModuleLeader；先限定学年和学期。'),
('Q2','教师办公室完整位置','Staff、Office、Building、Campus','hasOffice、locatedInBuilding、locatedOnCampus；roomNumber、floorNumber。'),
('Q3','教师联系方式','AcademicStaff','emailAddress、telephoneNumber、staffProfileURL。'),
('Q4','教师所属组织和所教模块','Staff、Organisation、ModuleOffering、Module','memberOf；通过 taughtBy 反向查开课，再用 offeringOf 查模块。'),
('Q5','专业负责人及联系方法','Course、CourseLeadershipAppointment、AcademicStaff','hasLeadershipAppointment、appointedStaffMember；筛选有效任期。'),
('Q6','科学楼位置与地图','Building、Campus','locatedOnCampus；buildingCode、postalAddress、mapURL、经纬度。'),
('Q7','806 房间位置与类型','Room 及其子类、Building','使用 roomID 消歧；查 floorNumber、locatedInBuilding 和 rdf:type。'),
('Q8','房间设备和容量','Room、EquipmentProvision、EquipmentType','hasEquipmentProvision、equipmentType；quantity、capacity。'),
('Q9','图书馆有电脑的学习室及预约状态','LibraryBuilding、StudySpace、EquipmentProvision','列出电脑数量大于零的学习室；返回 isBookable 的 true、false 或未知及 bookingURL。'),
('Q10','一年级模块教室无障碍及电梯','CourseModule、ModuleOffering、TeachingSession、AccessibleRoom、Building','筛选 studyYear=1 及培养方案；关联场次房间；isStepFree 与 hasLift 分别回答。'),
('Q11','专业简介院系及学制','Course、Department','offeredBy；courseTitle、courseDescription、qualificationName、durationYears。'),
('Q12','专业一年级模块必修或选修','Course、CourseModule、Module','hasCourseModule、forModule；studyYear、isCompulsory、curriculumVersion。'),
('Q13','先修与共修模块','Module','hasDirectPrerequisite、hasPrerequisite、hasCorequisite；区分直接与间接。'),
('Q14','模块考核及权重','ModuleOffering、Assessment 及其子类','hasAssessment；weightingPercentage、dueDateTime、提交信息。'),
('Q15','某场次教师地点和时间','ModuleOffering、TeachingSession、Room、AcademicStaff','hasTeachingSession、ledBy、takesPlaceIn、startDateTime、endDateTime。'),
('Q16 至 Q20','校园活动','CampusEvent、Organisation、Place','保留 organisedBy、takesPlaceIn、时间及报名字段；具体问题确定后再验收。'),
('Q21','设施要求与实际排课匹配','Module、TeachingSession、FacilityRequirement、FacilityType、Room','显式关联适用要求；比较 acceptableFacilityType 与 providesFacilityType。'),
('Q22','预计座位利用率偏低','TeachingSession、Room、ModuleOffering','按选定时段汇总预计人数和容量；排除取消场次并声明低利用率阈值。'),
]

CLASSES = [row for _, rows in CLASS_GROUPS for row in rows]
OBJECTS = [row for _, rows in OBJECT_GROUPS for row in rows]
DATA = [row for _, rows in DATA_GROUPS for row in rows]
SCHEMA = {'classes':CLASSES,'object_properties':OBJECTS,'data_properties':DATA,'annotations':ANNOTATIONS,'requirements':MAPPING}

def validate_schema():
    names = {r[0] for r in CLASSES}
    assert len(names) == len(CLASSES)
    for row in CLASSES:
        assert row[2] in names | {'owl:Thing'}, row
    for name, domain, range_, _ in OBJECTS:
        assert all(part in names for part in domain.split(' or ')), (name,domain)
        assert range_ in names, (name,range_)
    for name, domain, dtype, _ in DATA:
        assert domain in names | {'owl:Thing'}, (name,domain)
        assert dtype.startswith('xsd:'), name
    props = [r[0] for r in OBJECTS + DATA]
    assert len(props)==len(set(props))
    for i in range(1,16):
        assert any(r[0]==f'Q{i}' for r in MAPPING)

def set_font(style, size=11, bold=False, east='Arial Unicode MS'):
    style.font.name='Arial'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts'); rp.insert(0,rf)
    for key,value in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia',east)]: rf.set(qn('w:'+key),value)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+key),None)

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
sec.footer_distance=Inches(.25)
for name,size,bold in [('Normal',11,False),('Title',23,True),('Subtitle',11,False),('Heading 1',16,True),('Heading 2',12,True),('Heading 3',11,True)]:
    set_font(doc.styles[name],size,bold)
normal=doc.styles['Normal'].paragraph_format
normal.space_after=Pt(6); normal.line_spacing=1.15
normal.keep_with_next=False
normal.keep_together=False
normal.widow_control=True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[name].paragraph_format
    pf.space_before=Pt(12);pf.space_after=Pt(7);pf.keep_with_next=True
for name in ['Title','Subtitle']:
    doc.styles[name].paragraph_format.space_after=Pt(8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('G5 Team Hinton    ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='大学新生本体类与属性设计'
doc.core_properties.subject='类层级 属性字典 公理和需求映射'
doc.core_properties.author='G5 Team Hinton'
doc.core_properties.keywords='Ontology, OWL, Protégé, New Students'

def p(text='',style=None):
    return doc.add_paragraph(text,style)

NEXT_PAGE = False
def h(text,level=1):
    global NEXT_PAGE
    heading=doc.add_heading(text,level)
    if NEXT_PAGE and text.startswith(('2 ', '8 ')):
        heading.paragraph_format.page_break_before=True
    NEXT_PAGE=False
    return heading

def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',65),('bottom',65),('left',90),('right',90)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index,values in enumerate([headers]+list(rows)):
        row=t.rows[0] if index==0 else t.add_row()
        rowpr=row._tr.get_or_add_trPr();rowpr.append(OxmlElement('w:cantSplit'))
        for j,(cell,text) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(text)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE7F0' if index==0 else ('F5F7F9' if index%2==0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.line_spacing=1.08
                para.paragraph_format.keep_with_next=(index<=1)
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(10)
                    run.font.bold=(index==0)
                    run.font.color.rgb=RGBColor(0,0,0)
    spacer=p();spacer.paragraph_format.space_after=Pt(3);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=0.3
    return t

def code(text):
    para=p()
    para.paragraph_format.space_after=Pt(8)
    para.paragraph_format.line_spacing=1.05
    para.paragraph_format.keep_together=True
    run=para.add_run(text);run.font.name='Courier New';run.font.size=Pt(9.5)
    rf=run._element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Arial Unicode MS')
    return para

def page():
    global NEXT_PAGE
    NEXT_PAGE=True

validate_schema()
p('大学新生本体类与属性设计','Title')
p('G5 Team Hinton    2026年10月26日    版本 1.0','Subtitle')
p('本体面向本科与研究生新生，覆盖专业结构、教学人员、校园空间、教学安排和考核。本文给出可用于 Protégé 建库的类与属性字典，并把设施匹配和座位利用率作为扩展场景。')
p('核心设计区分专业 Course、模块 Module、具体开课 ModuleOffering 和教学场次 TeachingSession。必修属性保存在专业模块安排中，负责人任期和房间设备数量由独立记录承载，使跨专业、跨学年查询保持明确。')
h('1 范围与命名约定')
table(['项目','约定'],[
('需求范围','Q1 至 Q15 为基础问题；Q21 为设施匹配，Q22 为预计座位利用率。Q16 至 Q20 的活动需求另行细化。'),
('词汇规模',f'{len(CLASSES)} 个类，{len(OBJECTS)} 个对象属性，{len(DATA)} 个数据属性，{len(ANNOTATIONS)} 个注解属性。其中 FacilityRequirement 与 FacilityType 为设施扩展，CampusEvent 为活动接口。'),
('默认前缀','建议 ns: 对应 https://example.org/hinton/newstudent#。这是示例命名空间；正式建库前统一确定并保持稳定。'),
('命名规则','类使用单数 PascalCase；属性使用 camelCase；个体使用可区分年份、建筑或场次的标识。'),
('标签与说明','每个术语加中文和英文 rdfs:label，并用 rdfs:comment 说明定义与边界。'),
('专业名称','统一使用 Course。旧资料中的 Programme 按同一业务概念处理，不重复创建两个独立专业体系。'),
],[1.05,6.15])
h('概念区别',2)
p('类表示一类对象，例如 Room；个体表示具体对象，例如 ScienceBuilding_Room806；对象属性连接两个个体；数据属性连接个体与字面量。子类只表达“是一种”，空间包含与组织隶属通过对象属性表达。')
p('阅读顺序为类字典、对象属性、数据与注解属性、OWL 公理、数据验收、问题映射及示例。表中的父类、Domain 和 Range 均使用可直接在本体中创建的英文名称。')
page()

h('2 类与子类字典')
p('直接父类列给出类层级。除显式列出的定义类外，先建立命名类及其 SubClass Of 关系。顶层类的父类为 owl:Thing；同级类不自动互斥。')
for title,rows in CLASS_GROUPS:
    h(title,2)
    table(['Class 与中文名','直接父类','定义与边界'],[(f'{n}\n{zh}',parent,desc) for n,zh,parent,desc in rows],[2.32,1.75,3.13])
p('人员身份与房间用途可以重叠。例如研究生可以同时担任教学人员，研讨室也可以用作学习室。LibraryBuilding 仅表示物理建筑；需要描述图书馆管理部门时，应另建 Organisation 个体。')
p('StudySpace 在本版限定为房间。若后续加入大厅学习区或单个座位，应重新界定其父类和粒度。NewStudent 与具体入学记录、专业和参考学年有关，本版公共信息查询不将其建为永久身份。')
page()

h('3 对象属性字典')
p('对象属性的值是个体 IRI。Domain 和 Range 是推理公理，并非数据库字段校验。下表中出现 or 时，应在 Protégé 中录入一个并集表达式；分别添加多个 Domain 会表示交集。')
for title,rows in OBJECT_GROUPS:
    h(title,2)
    table(['属性','Domain 与 Range','含义与特性'],[(n,f'{d}\n→ {r}',desc) for n,d,r,desc in rows],[1.85,2.65,2.70])
p('组织结构只使用 organisationPartOf；空间位置使用 locatedInBuilding 与 locatedOnCampus。教师与负责人连接到 ModuleOffering，具体场次教师使用 ledBy。不要把整个教学团队自动复制为每次上课的教师。')
page()

h('4 数据属性与注解属性')
p('数据属性的值为文字、数字、日期或布尔值。以下表格统一列出 Domain 和 XSD 类型。未提供的值保持未知，不以 0 或 false 代替。字段是否必填由第 6 节的数据验收规则规定。')
for title,rows in DATA_GROUPS:
    h(title,2)
    table(['属性','Domain 与数据类型','定义与取值'],[(n,f'{d}\n{typ}',desc) for n,d,typ,desc in rows],[1.90,2.35,2.95])
h('注解属性',2)
table(['属性','适用对象','用途'],ANNOTATIONS,[1.80,1.30,4.10])
p('如果同一专业的全日制与非全日制路径有不同学制，可将其建成两个明确的 Course 变体。若必须共享专业身份再区分路径，应增加带修读模式和学制的路径记录；不要在同一个体上放多个无法配对的 studyMode 与 durationYears。')
p('成绩、学费、奖学金和学生个人注册记录不属于当前问答范围。需要新增这些问题时，再增加对应类和属性。')
page()

h('5 OWL 公理与推理设计')
p('以下示例采用 Manchester Syntax。基数限制描述模型中应存在的关系，不等同于已录入数据必须包含某字段。')
h('定义类',2)
code('Class: AccessibleRoom\n  EquivalentTo:\n    Room and (isStepFree value true)')
p('该定义仅表达无台阶可达，不代表具备所有无障碍设施。hasLift 或 accessibleEntrance 单独为 true 都不足以自动推得每个房间属于 AccessibleRoom。')
h('核心基数与值域限制',2)
table(['适用类','建议的 SubClass Of 表达式'],[
('ModuleOffering','offeringOf exactly 1 Module\ninAcademicYear exactly 1 AcademicYear\ninSemester min 1 Semester'),
('CourseModule','forCourse exactly 1 Course\nforModule exactly 1 Module\ninAcademicYear exactly 1 AcademicYear\nstudyYear exactly 1 xsd:positiveInteger\nisCompulsory exactly 1 xsd:boolean'),
('CourseLeadershipAppointment','appointedStaffMember exactly 1 AcademicStaff\ninverse hasLeadershipAppointment exactly 1 Course'),
('EquipmentProvision','equipmentType exactly 1 EquipmentType\nquantity exactly 1 xsd:nonNegativeInteger\ninverse hasEquipmentProvision exactly 1 Room'),
('Room','locatedInBuilding exactly 1 Building'),
('Building','locatedOnCampus exactly 1 Campus'),
('Semester','inAcademicYear exactly 1 AcademicYear'),
('OnCampusTeachingSession','TeachingSession\ntakesPlaceIn exactly 1 Room\ntakesPlaceIn only Room'),
('FacilityRequirement','acceptableFacilityType min 1 FacilityType'),
],[2.05,5.15])
p('Room、Building 的唯一直接位置以单个物理房间、单栋建筑的粒度为前提。一个 ModuleOffering 按一个学年建模，全年模块可连接多个 Semester。模块可联合负责、场次可多人授课、合班场次可由多个开课引用，因此这些关系不统一限制为 exactly 1。')
h('属性关系与位置推理',2)
code('ObjectProperty: hasCourseModule\n  InverseOf: forCourse\n\nObjectProperty: hasDirectPrerequisite\n  SubPropertyOf: hasPrerequisite\n\nObjectProperty: hasPrerequisite\n  Characteristics: Transitive\n\nObjectProperty: hasStudySpace\n  SubPropertyOf: inverse locatedInBuilding\n\nObjectProperty: withinCampus\n  SubPropertyChain: locatedInBuilding o locatedOnCampus')
p('最后一项表示属性链：房间位于某建筑、建筑位于某校区，可推出房间 withinCampus 该校区。它应在属性链公理中录入，而非把整行当成一个普通属性名称。withinCampus 和传递属性 hasPrerequisite 不用于 OWL 2 DL 的基数限制。')
p('先修关系中，多条直接先修默认表示每项均需满足。若课程规定“A 或 B”，应增加先修要求记录来表达替代条件。共修关系按学校规则确定方向，不默认设为 Symmetric。先修循环由查询检查。')
h('不相交类',2)
p('建议首先声明 Person、Organisation、Place 彼此不相交；Campus、Building、Room 彼此不相交；Course、Module、ModuleOffering 彼此不相交。其余组合在业务定义明确后再添加。')
p('Student 与 Staff 不设为不相交；Room 的用途子类不统一互斥；OnCampusTeachingSession 可同时属于 LectureSession、TutorialSession 或 PracticalSession。考试与课程作业是否互斥应由学校的考核分类定义决定。')
page()

h('6 数据验收与查询规则')
p('OWL 按开放世界解释数据：没有记录某属性，并不意味着该属性不存在或为假。对字段完整性、数值范围和跨记录一致性，应使用 SHACL、SPARQL 或应用检查。数据不足时返回“信息不足”，不要输出否定答案。')
table(['检查对象','验收规则','实现方式'],[
('身份与引用','实体有稳定 IRI；roomID 不重复；同楼层门牌号按建筑消歧。编号是否唯一不由 Functional 自动保证。','SHACL 或查询'),
('核心关联','开课、培养条目、任职与设备配置具有第 5 节所需的显式关联和值。','SHACL minCount 与 maxCount'),
('时间与学期','活动结束晚于开始；学期属于开课所选学年；任期结束不早于开始。日期时间采用一致时区规则。','SHACL SPARQL 或查询'),
('考核','完整公布的开课考核方案中，每项权重在 0 至 100，权重合计为 100。尚未公布完整时标记未完整。','SPARQL 聚合或应用'),
('房间设备','设备数量和容量为非负整数；查询“有电脑”还要求电脑配置 quantity 大于零。','SHACL 与查询'),
('预约','可预约使用 isBookable，必须预约使用 bookingRequired；没有 bookingURL 不等于不可预约。','查询与数据核实'),
('专业模块安排','查询同时指定专业、适用学年和培养方案版本；一年级用 studyYear=1。','SPARQL 条件'),
('来源与更新','维护来源 IRI 和 lastVerifiedOn；对任期、办公室、设备及课表更新进行核实。','维护流程'),
],[1.20,4.80,1.20])
h('设施匹配规则',2)
p('Q21 使用 FacilityRequirement 记录一项要求，并通过 acceptableFacilityType 列出可接受的类型。同一记录的多个类型表示任选其一；同一场次连接的多条要求表示每条都要满足。实际房间通过 providesFacilityType 记录已核实的设施类别。')
p('模块可声明设施要求，需检查的场次再显式关联相应要求记录。这样能区分讲授课和实验课，避免把某模块的实验设施要求自动施加给所有场次。设备种类与设施类别不同；如需由设备推得设施类型，必须给出明确映射规则。')
p('验收结果分为满足、不满足、信息不足。只有实际房间及设施清单足够完整时，才能判定不满足；尚未关联适用场次、缺少房间或设施信息时，应返回信息不足。')
h('预计座位利用率规则',2)
p('Q22 使用每场次 expectedStudentCount 和该场次房间 capacity。按场次等权汇总时，先逐场次计算“预计人数 ÷ 房间容量”，再取平均；如需按时长加权，则以“场次时长 × 该场次比例”为权重项求加权平均。统计前声明日期区间、纳入场次和低利用率阈值。')
p('对同一房间、容量固定且场次等时长的情况，可用“各场次预计人数之和 ÷ 房间容量 ÷ 场次数”得到相同结果。容量缺失或为零、预计人数缺失时不计算该记录，并报告遗漏数量。取消场次排除；合班场次按唯一场次去重，人数按该场次总人数记录。')
p('预计座位利用率不等于实际出勤利用率，也不等于房间时间利用率。计算实际出勤需要到课人数；计算时间利用率还需要房间可用时段。若统计跨越容量或设备配置变动，应增加有效期间记录或使用对应日期的快照。')
page()

h('7 问题到本体的映射')
p('下表按问题编号映射到类和查询路径。Q16 至 Q20 是活动需求接口，其余 17 个问题有明确场景。映射表说明模型如何支持问题，不代表查询已实现或已通过真实数据验收。')
table(['编号与主题','相关类','关键路径或条件'],[(f'{qid}\n{title}',classes,path) for qid,title,classes,path in MAPPING],[1.50,2.30,3.40])
page()

h('8 关联记录与个体示例')
p('以下为建模示例，名称、日期、数量和人员关联均用于演示，不代表学校实际安排。所有属性名均与本文字典一致。')
h('专业模块与开课',2)
code('BScComputerScience a UndergraduateCourse\n  hasCourseModule CS_Programming_Y1_2026\n\nCS_Programming_Y1_2026 a CourseModule\n  forCourse BScComputerScience\n  forModule IntroductionToProgramming\n  inAcademicYear AY2026_27\n  studyYear 1\n  isCompulsory true\n  curriculumVersion "2026"\n\nIntroductionToProgramming a Module\n\nProgramming_2026_Autumn a ModuleOffering\n  offeringOf IntroductionToProgramming\n  inAcademicYear AY2026_27\n  inSemester Autumn2026\n  taughtBy DrXiangLi\n  hasModuleLeader DrXiangLi\n  hasTeachingSession Programming_Lecture_01')
h('教学场次与空间',2)
code('Programming_Lecture_01\n  a LectureSession, OnCampusTeachingSession\n  ledBy DrXiangLi\n  takesPlaceIn ScienceBuilding_Room806\n  startDateTime "2026-10-26T09:00:00+00:00"\n  endDateTime "2026-10-26T10:00:00+00:00"\n  expectedStudentCount 24\n\nScienceBuilding_Room806 a SeminarRoom\n  locatedInBuilding ScienceBuilding\n  roomID "SCI-806"\n  roomNumber "806"\n  floorNumber "8"\n  capacity 40\n  isStepFree true\n  hasEquipmentProvision SCI806_Computers\n\nScienceBuilding a Building\n  locatedOnCampus MainCampus\n\nSCI806_Computers a EquipmentProvision\n  equipmentType ComputerEquipmentType\n  quantity 30\n\nComputerEquipmentType a EquipmentType')
p('这是便于阅读的事实清单，不是可直接导入的 Turtle 文件。正式建库时须为所有引用个体补齐类型，并为日期时间、整数和布尔值指定表中 XSD 类型。上述关系可推出 SCI-806 所在校区；isStepFree 为 true 可推出该房间属于 AccessibleRoom。')
h('任职记录',2)
code('BScComputerScience\n  hasLeadershipAppointment CS_Leader_2026\n\nCS_Leader_2026 a CourseLeadershipAppointment\n  appointedStaffMember DrXiangLi\n  validFrom "2026-08-01"\n  validUntil "2027-07-31"')
p('按查询日期筛选任职记录，再取得 appointedStaffMember 的联系方式。更换负责人时创建新任职记录，保留旧记录用于历史查询。')
page()

h('9 建库顺序与设计边界')
for text in [
'1 统一命名空间与术语，并按第 2 节从父类到子类创建 Classes。',
'2 创建对象属性，录入 Domain、Range、逆属性、子属性和属性链。',
'3 创建数据属性及 XSD 类型，补中文英文标签、注释与来源注解。',
'4 添加已明确的定义类、不相交公理和核心基数限制。',
'5 录入少量相互关联的个体，先覆盖专业、开课、教师、场次、房间和设备。',
'6 使用推理器检查一致性及预期推断，再用数据检查验证字段完整性。',
'7 按第 7 节逐题编写查询，记录输入条件、预期答案和实际结果。',
]:p(text)
h('需保持明确的范围',2)
p('活动子类随 Q16 至 Q20 的问题确定后再扩展。面向个人的选课、注册和新生身份判断，需要额外的学生注册记录及参考学年；它们不影响当前公共校园信息查询的类与属性。')
p('Q9 在表格中的中英文地点不一致，本文按英文 Central Library 采用图书馆场景。Q21 的设施匹配采用“每条要求满足一个可接受类型”的规则。Q22 采用预计座位利用率，阈值由查询输入指定。')
p('本方案在通用大学本体的人员、组织、空间和时间结构上加入开课、专业模块安排、任职与设备配置记录。它保留当前问答所需信息，并将年度变化与专业差异放在对应记录上。')
h('参考资料',2)
references=[
'[1] G5_Team_Hinton_questions.xlsx。“总体问题”B5 至 B10；“细节问题与分工”A4 至 E25。需求范围及问题编号。',
'[2] University本体_Protégé建库步骤.md。Step 1 至 Step 7。类、属性与公理的教学骨架。',
'[3] Day7_课堂分析.md。第六节与第七节。需求映射、子类语义及多个 Domain 的并集处理。',
'[4] Week_5_Reasoning_Pre_Lecture_Overview_Finalyi.pdf，第 1 页。公理、推断和一致性检查。',
'[5] W3C. OWL 2 Web Ontology Language Primer. https://www.w3.org/TR/owl2-primer/',
'[6] W3C. Defining N-ary Relations on the Semantic Web. https://www.w3.org/TR/swbp-n-aryRelations/',
'[7] W3C. Shapes Constraint Language. https://www.w3.org/TR/shacl/',
'[8] W3C. OWL 2 Manchester Syntax. https://www.w3.org/TR/owl2-manchester-syntax/',
]
for ref in references:
    para=p(ref);para.paragraph_format.space_after=Pt(5)
    for run in para.runs:run.font.size=Pt(10)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
(ROOT/'schema.json').write_text(json.dumps(SCHEMA,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUTPUT),'classes':len(CLASSES),'object_properties':len(OBJECTS),'data_properties':len(DATA),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
