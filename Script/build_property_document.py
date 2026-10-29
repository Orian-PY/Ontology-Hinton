from pathlib import Path
import sys,json,re
sys.path.insert(0,'/private/tmp/hinton-owl-check-deps')
from rdflib import Graph,Namespace,RDF,RDFS,OWL,URIRef
from rdflib.collection import Collection
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

WORK=Path('/private/tmp/hinton-place-properties')
OUTPUT=Path('/Users/peter/Ontology-Hinton/doc/Place属性设计.docx')
N=Namespace('https://example.org/hinton/newstudent#')
g=Graph().parse(WORK/'campus_academic_ontology.owl')
audit=json.loads((WORK/'audit.json').read_text())
report=json.loads((WORK/'validation.json').read_text())
assert report['result']=='passed'
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65)
section.left_margin=section.right_margin=Inches(.65)
section.header_distance=section.footer_distance=Inches(.28)

def font(st,size,bold=False):
    st.font.name='Arial';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
    pr=st.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.insert(0,rf)
    for key,val in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Arial Unicode MS')]:rf.set(qn('w:'+key),val)
for name,size,bold in [('Normal',11,False),('Title',23,True),('Subtitle',10.5,False),('Heading 1',15,True),('Heading 2',12,True),('Caption',9.5,False)]:font(doc.styles[name],size,bold)
for style in doc.styles:
    for border in list(style.element.findall('.//'+qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].paragraph_format.line_spacing=1.17
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name in ['Heading 1','Heading 2']:
    pf=doc.styles[name].paragraph_format;pf.space_before=Pt(12);pf.space_after=Pt(6);pf.keep_with_next=True
font(doc.styles['Header'],8.5);font(doc.styles['Footer'],8.5)
header=section.header.paragraphs[0];header.text='G5 Team Hinton    Place 属性设计';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
footer.add_run(' 页')

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text,level=1):return doc.add_heading(text,level)
def code(text):
    para=p(text)
    para.paragraph_format.space_before=Pt(3);para.paragraph_format.space_after=Pt(8)
    para.paragraph_format.line_spacing=1.08
    for run in para.runs:
        run.font.name='Courier New';run.font.size=Pt(9.5)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    return para

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Inches(width)
    tpr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        v=OxmlElement('w:'+edge);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'5');v.set(qn('w:color'),'D9D9D9');borders.append(v)
    tpr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,width in [('top',80),('bottom',80),('left',95),('right',95)]:
        v=OxmlElement('w:'+edge);v.set(qn('w:w'),str(width));v.set(qn('w:type'),'dxa');margins.append(v)
    tpr.append(margins)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for j,(cell,text,width) in enumerate(zip(cells,row,widths)):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=text
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6EF' if i==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.08
                para.paragraph_format.keep_together=True
                for run in para.runs:
                    run.font.size=Pt(10);run.font.bold=i==0
    p('').paragraph_format.space_after=Pt(1)
    return t

def local(term):
    text=str(term)
    return text.split('#')[-1] if '#' in text else text
def expr(term):
    if isinstance(term,URIRef):
        if str(term).startswith('http://www.w3.org/2001/XMLSchema#'):return 'xsd:'+local(term)
        return local(term)
    return ' or '.join(expr(v) for head in g.objects(term,OWL.unionOf) for v in Collection(g,head))
def property_type(name):
    return '\n'.join(expr(d) for d in g.objects(N[name],RDFS.domain))+'\n→ '+ '\n'.join(expr(r) for r in g.objects(N[name],RDFS.range))
def status(name):
    if name in audit['new_data_properties']+audit['new_object_properties']:return '新增'
    if name in audit['modified_existing_entities']:return '调整'
    return '保留'

existing_zh={
    'locatedInBuilding':'房间直接位于的建筑；保留每个物理房间恰好一栋建筑的限制。',
    'locatedOnCampus':'建筑或室外空间直接所属的校区；Domain 使用一个并集。',
    'spatiallyPartOf':'直接或间接的空间包含，具有传递性；与组织隶属区分。',
    'withinCampus':'由房间所在建筑及建筑所在校区自动推出；也属于 spatiallyPartOf。',
    'hasLibraryRoom':'图书馆建筑包含的图书馆房间；是 hasRoom 的子属性。',
    'hasStudySpace':'图书馆建筑包含的学习室；是 hasRoom 的子属性。',
    'hasOffice':'保留人员部分已有的办公地点关联，允许多个办公室。',
    'hasTeachingRoom':'保留模块部分已有的教室关联；没有具体课次或时间配对信息。',
    'postalAddress':'适用地点的邮寄地址，Domain 从 Building 扩为 Place。',
    'mapURL':'地点官方地图链接，Domain 从 Building 扩为 Place。',
    'latitude':'代表地点位置的纬度，范围 −90 至 90，Domain 扩为 Place。',
    'longitude':'代表地点位置的经度，范围 −180 至 180，Domain 扩为 Place。',
    'areaSquareMetres':'地点面积，单位平方米；非负；允许用于房间和室外空间。',
    'openingHoursText':'面向读者的开放时段文字；不能直接用于计算空闲时段。',
    'isBookable':'是否支持提前预约；true、false 与未提供需分别处理。',
    'bookingRequired':'是否必须预约；与是否支持预约分别记录。',
    'bookingURL':'预约入口；没有链接不代表不可预约，也不表示实时空闲。',
    'campusName':'校区名称；是 placeName 的子属性。',
    'campusAddress':'校区地址；是 postalAddress 的子属性。',
    'campusMapURL':'校区地图；是 mapURL 的子属性。',
    'buildingName':'建筑名称；是 placeName 的子属性。',
    'buildingCode':'建筑代码；是 placeCode 的子属性。',
    'hasLift':'建筑是否有电梯；不能单独推断某个房间可以无障碍进入。',
    'accessibleEntrance':'建筑是否至少有一个已核实的无台阶入口。',
    'libraryServiceURL':'图书馆服务页；是 serviceURL 的子属性。',
    'spaceName':'室外空间名称；是 placeName 的子属性。',
    'spaceCode':'室外空间编码；是 placeCode 的子属性。',
    'surfaceType':'房间或室外空间的地面材质；Domain 扩为二者的一个并集。',
    'roomID':'跨建筑消歧的房间标识；是 placeCode 的子属性，唯一性另行验收。',
    'roomName':'房间名称；是 placeName 的子属性。',
    'roomNumber':'门牌号，例如 806；不同建筑内可以重复。',
    'floorNumber':'楼层字符串，例如 G、B1、8，不强制为整数。',
    'capacity':'房间允许容纳人数；计算座位利用率时须大于零。',
    'isStepFree':'是否已核实存在通往房间的无台阶路线；缺失为未知。',
    'isAccessible':'轮椅可经连续无台阶路线到达，并可通过适合轮椅的入口进入。',
    'accessibilityNote':'入口、路线、电梯、门口限制及核实依据；不等于全面认证。',
    'hasHearingLoop':'是否已核实配有助听感应环；缺失为未知。',
    'safetyInstructions':'实验室安全使用要求的文字说明。',
    'isQuietStudy':'学习室是否被指定为安静学习区域。',
    'openingHours':'旧的建筑或房间开放时段；作为 openingHoursText 的子属性保留。',
    'roomAreaSquareMetres':'旧的房间可用面积字段；作为 areaSquareMetres 的子属性保留。',
}
desc={**existing_zh,**{k:v[-1] for k,v in audit['object_new'].items()},**{k:v[-1] for k,v in audit['data_new'].items()}}
included=[]
def properties(names):
    rows=[]
    for name in names:
        assert list(g.objects(N[name],RDF.type)),name
        rows.append([name+'\n'+status(name),property_type(name),desc[name]])
        included.append(name)
    table(['Property 与状态','Domain 与 Range','用途与边界'],rows,[2.05,2.25,2.9])

p('Place 属性设计','Title')
p('大学新生本体    2026年10月28日    本体文件版本 2.8','Subtitle')
p('依据类设计 1.3 版与现有 OWL；新增字段为本次设计，取值须依据学校资料或核实结果。','Caption')
p('我们以类设计文档 1.3 版中的地点层级为基础，把名称、地图、预约和联系信息集中在 Place，把直接位置和房间能力保留在适合的父类。所有地点子类继承父类的适用属性，专有字段只补充真实存在的业务区别。')
p('本次新增 7 个对象属性、16 个数据属性，并采用 EquipmentType 和 EquipmentProvision 两个辅助类记录设备种类与数量。属性均为可逐步补录的模式定义，不要求每个地点填写全部字段。完整 OWL 共有 62 个命名类、30 个命名对象属性和 106 个命名数据属性。')
h('1 属性的使用约定')
p('对象属性连接两个个体，例如 Room806 locatedInBuilding ScienceBuilding。数据属性连接个体与字面量，例如 Room806 capacity 40。Domain 和 Range 会参与类型推断，因此使用 capacity 会把主体归入 Room；宿舍楼住宿人数另用 residentCapacity。')
p('表中“新增”表示本次加入；“调整”表示扩展适用范围或增加子属性关系；“保留”表示沿用已有定义。or 表示单个并集表达式，不能拆成多条 Domain。布尔字段缺失时按未知处理。')
h('2 对象属性')
h('空间位置与包含',2)
properties(['locatedInBuilding','hasRoom','locatedOnCampus','hasCampusPlace','withinCampus','spatiallyPartOf','containsPlace','hasLibraryRoom','hasStudySpace'])
h('管理设备与外部关联',2)
properties(['managedBy','hasEquipmentProvision','provisionForRoom','equipmentType','hasOffice','hasTeachingRoom'])
p('EquipmentType 和 EquipmentProvision 是独立辅助类，不属于 Place。每个配置记录只描述一间房间的一类设备及其数量。')
h('3 Place 共享数据属性')
p('下列属性适用于校区、建筑、室内房间和室外空间。地图、经纬度和邮寄地址只有在有适用信息时填写；位置关系不会自动复制建筑的地址或坐标到房间。')
properties(['placeName','placeCode','locationDescription','postalAddress','mapURL','latitude','longitude','areaSquareMetres','openingHoursText','isBookable','bookingRequired','bookingURL','accessPolicy','serviceURL','placeContactEmail','placeContactTelephone'])
h('4 各类地点的专有数据属性')
h('校区建筑与室外空间',2)
properties(['campusName','campusAddress','campusMapURL','buildingName','buildingCode','hasLift','accessibleEntrance','libraryServiceURL','spaceName','spaceCode','surfaceType','hasFloodlighting','supportedSport','residentCapacity','cateringInformation'])
p('SportsHall 继承 SportsBuilding 的 supportedSport；OutdoorCourt 和 SportsGround 继承 OutdoorSpace 的名称、地面和照明属性。TeachingBuilding、LaboratoryBuilding、AdministrationBuilding 等建筑主要复用 Building 和 Place 的公共字段；管理组织通过 managedBy 表达。')
h('室内房间教学与图书馆空间',2)
properties(['roomID','roomName','roomNumber','floorNumber','capacity','isStepFree','isAccessible','accessibilityNote','hasHearingLoop','hasFixedSeating','safetyInstructions','requiresSafetyInduction','isQuietStudy','collectionDescription','openingHours','roomAreaSquareMetres'])
p('AcademicOffice 和 AdministrativeOffice 继承 Office、Room 与 Place 的属性。LibraryReadingRoom 同时继承 ReadingRoom 和 LibraryRoom，因此可使用 StudySpace 的安静学习字段与 Room 的容量和无障碍字段。LoanServiceRoom 直接复用开放时段、服务页和联系字段。')
p('LectureTheatre、Classroom 和 Laboratory 通过 TeachingRoom 继承 hasFixedSeating。原有 SeminarRoom 本次保持原类层级，仍直接继承 Room；确认其属于教学房间时，可另外断言该房间也是 TeachingRoom。体育房间可以使用容量和地面材质，无需为每种球类再建立一套相同属性。')
h('设备配置的数据字段',2)
properties(['equipmentName','quantity'])
h('5 子属性与推理关系')
p('录入专用属性后，可以在启用推理的查询中统一检索通用属性。方向是从子属性推出父属性，填写 placeName 不会自动生成 buildingName 或 roomName。')
code('campusName, buildingName, roomName, spaceName\n  SubPropertyOf placeName\nbuildingCode, roomID, spaceCode\n  SubPropertyOf placeCode\ncampusAddress SubPropertyOf postalAddress\ncampusMapURL SubPropertyOf mapURL\nlibraryServiceURL SubPropertyOf serviceURL\nopeningHours SubPropertyOf openingHoursText\nroomAreaSquareMetres SubPropertyOf areaSquareMetres')
code('hasRoom InverseOf locatedInBuilding\nhasCampusPlace InverseOf locatedOnCampus\ncontainsPlace InverseOf spatiallyPartOf\nprovisionForRoom InverseOf hasEquipmentProvision\nhasLibraryRoom, hasStudySpace SubPropertyOf hasRoom\nlocatedInBuilding o locatedOnCampus -> withinCampus')
p('spatiallyPartOf 与 containsPlace 具有传递性，查询结果可能包含间接包含的地点。若要列出建筑直接包含的房间，使用 hasRoom；若要列出校区直接包含的建筑或室外空间，使用 hasCampusPlace。传递属性和属性链推导的 withinCampus 不用于基数限制。')
h('6 约束与数据验收')
table(['适用类','OWL 约束'],[
    ['EquipmentProvision','equipmentType exactly 1 EquipmentType\nprovisionForRoom exactly 1 Room\nquantity exactly 1 xsd:nonNegativeInteger'],
    ['TeachingRoom','hasFixedSeating max 1 xsd:boolean'],
    ['Laboratory','requiresSafetyInduction max 1 xsd:boolean'],
    ['OutdoorSpace','hasFloodlighting max 1 xsd:boolean'],
    ['StudentResidence','residentCapacity max 1 xsd:nonNegativeInteger'],
],[1.7,5.5])
p('保留既有的房间和建筑唯一直接位置、OutdoorSpace 唯一校区，以及 Room 的 isAccessible、isStepFree 至多一个值等约束。OWL 的 exactly 1 允许尚未命名的值存在，不能代替“录入必须完整”的检查；填写完整性、编号唯一性和数值范围需用 SHACL、查询或应用另行验收。')
p('建议验收经纬度上下界、面积非负、用于座位计算的 capacity 大于零。hasLift 和 accessibleEntrance 不能单独推出某个房间的 isAccessible。requiresSafetyInduction 说明使用要求，不记录个人是否通过培训。布尔值 true、false 与未知均应保留。')
p('来源可用 dcterms:source 保存资料 IRI。原文件的 lastVerifiedOn 采用 xsd:dateTime，以带时区的零点表示核实日期；本次沿用此编码。开放时段、预约链接和照明信息分别描述各自事实，不自动表示当前空闲或夜间开放。')
h('7 示例与可支持的问题')
p('以下是示意性的 Turtle 事实片段，只用于说明字段如何配对，不代表学校的真实地点或设备。示例不写入交付的 OWL。')
code('@prefix ns: <https://example.org/hinton/newstudent#> .\n@prefix xsd: <http://www.w3.org/2001/XMLSchema#> .\n\nns:LibraryRoom01 a ns:ReadingRoom ;\n  ns:locatedInBuilding ns:CentralLibrary ;\n  ns:roomID "LIB-101" ;\n  ns:isBookable false ;\n  ns:hasEquipmentProvision ns:LIB101Computers .\nns:CentralLibrary a ns:LibraryBuilding .\nns:LIB101Computers a ns:EquipmentProvision ;\n  ns:equipmentType ns:ComputerType ;\n  ns:quantity "12"^^xsd:nonNegativeInteger .\nns:ComputerType a ns:EquipmentType ;\n  ns:equipmentName "Computer" .')
p('该位置可推断为 LibraryRoom 和 LibraryReadingRoom；设备配置可以查出电脑数量为 12。Q9 查询“有电脑的图书馆学习室”时，要求对应电脑配置 quantity 大于零，再返回 isBookable 的 true、false 或未知；不能先筛选 isBookable=true，否则会漏掉不可预约或预约状态未知的房间。')

# Cross-check the dictionary with the actual deliverable before saving.
assert len(included)==len(set(included))
assert len(included)==64,len(included)
assert set(audit['new_object_properties']+audit['new_data_properties'])<=set(included)
doc.core_properties.title='Place 属性设计'
doc.core_properties.subject='大学新生本体地点属性字典'
doc.core_properties.author='G5 Team Hinton'
doc.save(OUTPUT)
(WORK/'document-audit.json').write_text(json.dumps({'output':str(OUTPUT),'properties_documented':len(included),'object_properties':15,'data_properties':49,'ontology_version':'2.8','new_properties_covered':True},ensure_ascii=False,indent=2))
print(OUTPUT)
